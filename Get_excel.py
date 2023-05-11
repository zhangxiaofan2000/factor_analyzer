# -*- coding: utf-8 -*-
# Auther : ZhangYiLong
# Mail : 503302425@qq.com
# Date : 2023/5/11 23:11
# File : Get_excel.py
import os
import pandas as pd
import docx
import concurrent.futures
from tqdm import tqdm


folder_path = r'.\data\word'
output_dir = r'.\data\财务excel'


def get_table_data(table):
    df = pd.DataFrame()
    try:
        if '营业收入' in table.cell(1, 0).text.strip():
            rows = table.rows
            header = [cell.text for cell in rows[0].cells]
            data = [[cell.text for cell in row.cells] for row in rows[1:]]
            df = pd.DataFrame(data, columns=header,index=None)
    except:
        pass
    return df.iloc[:, :2].T # 只取前两列数据

def extract_table_data(filename):
    # 打开word文档
    doc = docx.Document(folder_path+"\\"+filename)

    # 遍历文档中的所有表格
    for i, table in enumerate(doc.tables):
        # 获取表格数据
        df = get_table_data(table)
        if not df.empty:
            # 保存表格数据到excel中
            output_file = os.path.join(output_dir, f"{filename}.xlsx")
            df.to_excel(output_file, index=False, encoding="utf-8-sig")
            break


files = os.listdir(folder_path)

with concurrent.futures.ThreadPoolExecutor() as executor:
    # 提交任务到线程池中并获得Future对象
    future_list = [executor.submit(extract_table_data, arg) for arg in files]

    # 获取每个任务的结果
    result_list = []
    with tqdm(total=len(future_list)) as pbar:
        for future in concurrent.futures.as_completed(future_list):
            result = future.result()
            result_list.append(result)
            pbar.update(1)


    # 并发执行任务
    # results2 = list(tqdm(executor.map(extract_table_data,files), total=len(files)))