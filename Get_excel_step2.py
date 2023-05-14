# -*- coding: utf-8 -*-
# Auther : ZhangYiLong
# Mail : 503302425@qq.com
# Date : 2023/5/11 23:11
# File : Get_excel.py
import asyncio
import os
import time

import pandas as pd
import docx
import concurrent.futures
from tqdm import tqdm


folder_path = r'.\data\word'
output_dir = r'.\data\资产excel'


def get_table_data(table):
    df = pd.DataFrame()
    try:
        # if '销售费用' in table.cell(1, 0).text.strip():
        if '货币资金' in table.cell(1, 0).text.strip():
            rows = table.rows
            header = [cell.text for cell in rows[0].cells]
            data = [[cell.text for cell in row.cells] for row in rows[1:]]
            df = pd.DataFrame(data, columns=header,index=None)
    except:
        pass
    return df.iloc[:, :2].T.reset_index(drop=False)


def extract_table_data(filename):
    try:
        # 打开word文档
        doc = docx.Document(folder_path+"\\"+filename)
        # 遍历文档中的所有表格
        for i, table in enumerate(doc.tables):
            # 获取表格数据
            df = get_table_data(table)
            if not df.empty:
                # 保存表格数据到excel中
                output_file = os.path.join(output_dir, f"{filename}.xlsx")
                df.to_excel(output_file, index=False, encoding="utf-8-sig")
                break
    except:
        print(filename)


if __name__ == '__main__':

    files = os.listdir(folder_path)

    with concurrent.futures.ProcessPoolExecutor() as executor:
        # 并发执行任务
        results = list(tqdm(executor.map(extract_table_data, files), total=len(files)))
# async def main(files):
#     loop = asyncio.get_running_loop()
#     with concurrent.futures.ThreadPoolExecutor() as pool:
#         # 创建一个线程池，用于执行IO操作
#         tasks = []
#         for file in files:
#             # 提交异步任务到线程池
#             task = loop.run_in_executor(pool, extract_table_data, file)
#             tasks.append(task)
#         # 并发执行任务
#         results = await asyncio.gather(*tasks)
#         return results

# # 调用异步函数
# files = os.listdir(folder_path)
# results = asyncio.run(main(files))
#
#
# with concurrent.futures.ThreadPoolExecutor() as executor:
#     # 提交任务到线程池中并获得Future对象
#     future_list = [executor.submit(extract_table_data, arg) for arg in files]
#
#     # 获取每个任务的结果
#     result_list = []
#     with tqdm(total=len(future_list)) as pbar:
#         for future in concurrent.futures.as_completed(future_list):
#             result = future.result()
#             result_list.append(result)
#             pbar.update(1)

